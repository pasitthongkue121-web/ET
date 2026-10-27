# -*- coding: utf-8 -*-
"""
Generate PowerPoint Presentation for Design Thinking Playbook
Compliant with Pitching Rubric & 7-minute presentation structure
"""
import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

def create_deck():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    # Theme Colors
    DARK_BG = RGBColor(15, 23, 42)      # Slate 900
    CARD_BG = RGBColor(30, 41, 59)      # Slate 800
    WHITE = RGBColor(255, 255, 255)
    LIGHT_GRAY = RGBColor(203, 213, 225) # Slate 300
    MUTED_GRAY = RGBColor(148, 163, 184) # Slate 400
    ACCENT_BLUE = RGBColor(56, 189, 248) # Sky 400
    ACCENT_CYAN = RGBColor(6, 182, 212)  # Cyan 500
    ACCENT_GREEN = RGBColor(52, 211, 153)# Emerald 400
    ACCENT_AMBER = RGBColor(251, 191, 36)# Amber 400
    ACCENT_PURPLE = RGBColor(167, 139, 250)# Purple 400
    BORDER_COLOR = RGBColor(51, 65, 85) # Slate 700

    def add_header(slide, step_tag, title, subtitle):
        # Background
        bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
        bg.fill.solid()
        bg.fill.fore_color.rgb = DARK_BG
        bg.line.fill.background()

        # Step Tag
        tag_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(11.733), Inches(0.4))
        tf_tag = tag_box.text_frame
        tf_tag.word_wrap = True
        p_tag = tf_tag.paragraphs[0]
        p_tag.text = step_tag.upper()
        p_tag.font.size = Pt(11)
        p_tag.font.bold = True
        p_tag.font.color.rgb = ACCENT_BLUE

        # Main Title
        title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.75), Inches(11.733), Inches(0.6))
        tf_title = title_box.text_frame
        tf_title.word_wrap = True
        p_title = tf_title.paragraphs[0]
        p_title.text = title
        p_title.font.size = Pt(24)
        p_title.font.bold = True
        p_title.font.color.rgb = WHITE

        # Subtitle
        sub_box = slide.shapes.add_textbox(Inches(0.8), Inches(1.35), Inches(11.733), Inches(0.4))
        tf_sub = sub_box.text_frame
        tf_sub.word_wrap = True
        p_sub = tf_sub.paragraphs[0]
        p_sub.text = subtitle
        p_sub.font.size = Pt(13)
        p_sub.font.color.rgb = MUTED_GRAY

    def add_card(slide, left, top, width, height, title, items, badge="", accent=ACCENT_BLUE):
        card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
        card.fill.solid()
        card.fill.fore_color.rgb = CARD_BG
        card.line.color.rgb = BORDER_COLOR
        card.line.width = Pt(1)

        tx_box = slide.shapes.add_textbox(left + Inches(0.2), top + Inches(0.2), width - Inches(0.4), height - Inches(0.4))
        tf = tx_box.text_frame
        tf.word_wrap = True

        if badge:
            p_badge = tf.paragraphs[0]
            p_badge.text = badge.upper()
            p_badge.font.size = Pt(10)
            p_badge.font.bold = True
            p_badge.font.color.rgb = accent
            p_title = tf.add_paragraph()
        else:
            p_title = tf.paragraphs[0]

        p_title.text = title
        p_title.font.size = Pt(16)
        p_title.font.bold = True
        p_title.font.color.rgb = WHITE
        p_title.space_after = Pt(10)

        for item in items:
            p = tf.add_paragraph()
            p.text = "• " + item
            p.font.size = Pt(12)
            p.font.color.rgb = LIGHT_GRAY
            p.space_after = Pt(6)

    # -------------------------------------------------------------
    # SLIDE 1: COVER
    # -------------------------------------------------------------
    s1 = prs.slides.add_slide(blank_layout)
    bg1 = s1.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
    bg1.fill.solid()
    bg1.fill.fore_color.rgb = DARK_BG
    bg1.line.fill.background()

    # Title box
    tbox = s1.shapes.add_textbox(Inches(1.0), Inches(1.5), Inches(11.333), Inches(4.5))
    tf1 = tbox.text_frame
    tf1.word_wrap = True

    p0 = tf1.paragraphs[0]
    p0.text = "DESIGN THINKING PLAYBOOK & PITCHING"
    p0.font.size = Pt(14)
    p0.font.bold = True
    p0.font.color.rgb = ACCENT_BLUE
    p0.space_after = Pt(10)

    p1 = tf1.add_paragraph()
    p1.text = "คู่มือปั้นนวัตกรรมฉบับนักศึกษา “จากไอเดียสู่เว็บแอปพลิเคชัน”"
    p1.font.size = Pt(32)
    p1.font.bold = True
    p1.font.color.rgb = WHITE
    p1.space_after = Pt(15)

    p2 = tf1.add_paragraph()
    p2.text = "กรณีศึกษา CampusRide — แพลตฟอร์มสองด้าน (Two-Sided Platform) เรียกรถรับส่งภายในมหาวิทยาลัย"
    p2.font.size = Pt(18)
    p2.font.color.rgb = LIGHT_GRAY
    p2.space_after = Pt(30)

    p3 = tf1.add_paragraph()
    p3.text = "วิชาแนวคิดแบบผู้ประกอบการ ประจำภาคการศึกษาที่ 1/2569\nคณะครุศาสตร์อุตสาหกรรมและเทคโนโลยี มหาวิทยาลัยเทคโนโลยีราชมงคลศรีวิชัย\nอาจารย์ที่ปรึกษา: ผศ.ดร.ไชยยะ ธนพัฒน์ศิริ"
    p3.font.size = Pt(13)
    p3.font.color.rgb = MUTED_GRAY

    # -------------------------------------------------------------
    # SLIDE 2: 3 PILLARS & 4 GOLDEN RULES
    # -------------------------------------------------------------
    s2 = prs.slides.add_slide(blank_layout)
    add_header(s2, "FOUNDATION · กรอบคิดนวัตกรรมระดับโลก", "สามเสาหลัก (Sweet Spot) & สี่กฎเหล็กของนักออกแบบ", "สมดุลของนวัตกรรมที่ยั่งยืน และกรอบความคิดแบบผู้ประกอบการเทคโนโลยี")
    
    col_w = Inches(3.64)
    top_pos = Inches(1.9)
    h_pos = Inches(2.5)

    add_card(s2, Inches(0.8), top_pos, col_w, h_pos, "1. Desirability", 
             ["หัวใจ: ความต้องการของมนุษย์", "คำถามหลัก: 'มีคนต้องการสิ่งนี้จริงหรือไม่?'", "ต้องเข้าใจ Pain Point และแก้ปัญหาได้ตรงจุด", "ตรวจสอบในขั้นตอน: Empathize, Define, Test"],
             "ความต้องการของคน", ACCENT_BLUE)
    
    add_card(s2, Inches(4.84), top_pos, col_w, h_pos, "2. Feasibility", 
             ["หัวใจ: ความเป็นไปได้ทางเทคนิค", "คำถามหลัก: 'เราสร้างได้จริงด้วยทรัพยากรที่มีไหม?'", "ใช้ Google Apps Script + Google Sheets + AI", "ตรวจสอบในขั้นตอน: Ideate, Prototype, Test"],
             "เทคโนโลยีและความเป็นไปได้", ACCENT_GREEN)

    add_card(s2, Inches(8.88), top_pos, col_w, h_pos, "3. Viability", 
             ["หัวใจ: ความอยู่รอดทางธุรกิจ", "คำถามหลัก: 'สิ่งนี้จะสร้างรายได้และยั่งยืนไหม?'", "ทดสอบด้วย Fake Door / Pre-order จริง", "ตรวจสอบในขั้นตอน: Test & Iterate"],
             "ความยั่งยืนทางการเงิน", ACCENT_AMBER)

    # 4 Rules Card below
    add_card(s2, Inches(0.8), Inches(4.6), Inches(11.733), Inches(2.3), "4 กฎเหล็กของนักออกแบบนวัตกรรม (Core Mindsets)", 
             [
                 "Iterative Process คือหัวใจ: ย้อนกลับไปทำซ้ำไม่ใช่ความล้มเหลว แต่คือ Learning Loop และการ Pivot",
                 "User กับ Customer อาจไม่ใช่คนเดียวกัน: แยกให้ออกระหว่างผู้ใช้งานแอปฯ กับผู้มีอำนาจจ่ายเงิน",
                 "Deploy Empathy ก่อน Solution: อย่าเพิ่งสร้างฟีเจอร์ถ้ายังไม่เข้าใจ Pain Point ที่แท้จริงของผู้ใช้",
                 "Fail Fast, Fail Cheap: ล้มเหลวให้เร็วและประหยัดที่สุด เพื่อลดต้นทุนความผิดพลาดในการสร้างจริง"
             ],
             "MINDSET", ACCENT_PURPLE)

    # -------------------------------------------------------------
    # SLIDE 3: STEP 1 - EMPATHIZE
    # -------------------------------------------------------------
    s3 = prs.slides.add_slide(blank_layout)
    add_header(s3, "ขั้นตอนที่ 1 · EMPATHIZE (นาทีที่ 0:00 - 1:00)", "เข้าใจ ให้ถึงใจ: สำรวจ Pain Point เชิงลึก", "เครื่องมือภาคสนาม AEIOU และเทคนิคขุดลึกสาเหตุรากเหง้าด้วย 5 Whys")

    add_card(s3, Inches(0.8), Inches(1.9), Inches(5.66), Inches(4.9), "กรอบการสังเกตการณ์ AEIOU", 
             [
                 "Activities (กิจกรรม): นักศึกษาเร่งรีบข้ามสะพานลอย หอบหนังสือและโน้ตบุ๊ก",
                 "Environment (สิ่งแวดล้อม): มหาวิทยาลัยมีพื้นที่กว้างมาก แดดร้อนจัดและฝนตกบ่อย",
                 "Interactions (ปฏิสัมพันธ์): ต่อแถวรอวินมอเตอร์ไซค์หน้าประตูยาวเหยียด วินขาดช่วงเร่งด่วน",
                 "Objects (สิ่งของ): สมาร์ทโฟนที่เปิดแอปฯ เช็คชื่อคาบเรียน, กระเป๋าหนังสือ, รถมอเตอร์ไซค์",
                 "Users (ผู้ใช้): นักศึกษาปี 1 ไม่ชำนาญเส้นทาง และนักศึกษารุ่นพี่ที่มีรถจักรยานยนต์ส่วนตัว",
                 "ข้อค้นพบสำคัญ: การเดินข้ามตึกใช้เวลา 20-30 นาที ทำให้เข้าเรียนสายเป็นประจำ"
             ],
             "CONTEXTUAL OBSERVATION", ACCENT_BLUE)

    add_card(s3, Inches(6.86), Inches(1.9), Inches(5.66), Inches(4.9), "เทคนิคขุดรากเหง้าด้วย 5 Whys", 
             [
                 "คำถาม 1: ทำไมถึงยอมจ่ายค่ามอเตอร์ไซค์รับจ้างแพงกว่าปกติ?",
                 "คำตอบ: เพราะกลัวเข้าห้องเรียนสายระหว่างเปลี่ยนคาบ",
                 "คำถาม 2: ทำไมถึงกลัวเข้าเรียนสาย?",
                 "คำตอบ: เพราะอาจารย์บางวิชาปิดประตูห้องและเช็คชื่อตรงเวลาเป๊ะ",
                 "คำถาม 3: ทำไมการเดินปกติถึงไปไม่ทัน?",
                 "คำตอบ: คาบเรียนเช้าและบ่ายอยู่คนละฟากมหาวิทยาลัย ระยะทางไกล",
                 "คำถาม 4: ทำไมไม่มีทางเลือกอื่น?",
                 "คำตอบ: มอเตอร์ไซค์รับจ้างคิวแน่น และค่าบริการสูงถึง 30-40 บาท/เที่ยว",
                 "Root Cause (สาเหตุแท้จริง): ขาดระบบขนส่งราคาประหยัดที่เชื่อมตึกเรียนตามตารางเวลาเรียน"
             ],
             "ROOT CAUSE ANALYSIS", ACCENT_CYAN)

    # -------------------------------------------------------------
    # SLIDE 4: STEP 2 - DEFINE (TWO-SIDED PERSONAS)
    # -------------------------------------------------------------
    s4 = prs.slides.add_slide(blank_layout)
    add_header(s4, "ขั้นตอนที่ 2 · DEFINE (นาทีที่ 1:00 - 2:00)", "สังเคราะห์ข้อมูล เพื่อตีโจทย์แพลตฟอร์มสองด้าน", "สร้าง Persona ตัวแทนผู้ใช้จริงทั้งฝั่งผู้รับบริการ (Consumer) และผู้ให้บริการ (Provider)")

    add_card(s4, Inches(0.8), Inches(1.9), Inches(5.66), Inches(3.7), "Persona: ฝั่งผู้รับบริการ (Consumer)", 
             [
                 "ชื่อ: น้องปาล์ม รีบไปเรียน (อายุ 19 ปี, นักศึกษาปี 1)",
                 "เป้าหมาย: เดินทางข้ามฝั่งมหาวิทยาลัยให้ทันคาบเรียนถัดไปโดยไม่ถูกเช็คชื่อสาย",
                 "Pain Point: มอเตอร์ไซค์รับจ้างแพงช่วงเร่งด่วน และต้องยืนรอคิวกลางแดดนาน",
                 "คุณค่าที่ต้องการ: เรียกรถได้ภายใน 5 นาที ราคาประหยัด เห็นตำแหน่งคนขับแบบเรียลไทม์",
                 "พฤติกรรม: ใช้สมาร์ทโฟนเป็นหลัก คุ้นเคยกับการกดเรียกรถแอปฯ ยอดนิยม"
             ],
             "CONSUMER PERSONA", ACCENT_BLUE)

    add_card(s4, Inches(6.86), Inches(1.9), Inches(5.66), Inches(3.7), "Persona: ฝั่งผู้ให้บริการ (Provider)", 
             [
                 "ชื่อ: พี่บอล อยากมีรายได้เสริม (อายุ 21 ปี, นักศึกษาปี 3)",
                 "เป้าหมาย: มีรถมอเตอร์ไซค์ส่วนตัว อยากหารายได้พิเศษช่วงว่างระหว่างคาบ",
                 "Pain Point: ไม่สามารถทำงานพิเศษประจำได้เพราะตารางเรียนไม่แน่นอน",
                 "ความกังวล: กังวลเรื่องความปลอดภัย และกลัวถูกโกงค่าโดยสารหรือไม่จ่ายเงิน",
                 "คุณค่าที่ต้องการ: เลือกเวลารับงานได้อิสระ และระบบยืนยันตัวตนคนนั่งด้วยอีเมลสถาบัน"
             ],
             "PROVIDER PERSONA", ACCENT_GREEN)

    # Problem Statement Card below
    add_card(s4, Inches(0.8), Inches(5.75), Inches(11.733), Inches(1.35), "คมคำนิยามปัญหา (Problem Statements)", 
             [
                 "ฝั่ง Consumer: 'น้องปาล์มต้องการวิธีเดินทางข้ามมหาวิทยาลัยที่รวดเร็วและราคาย่อมเยา เพราะกลัวเข้าเรียนสายและไม่อยากจ่ายแพงเกินจริง'",
                 "ฝั่ง Provider: 'พี่บอลต้องการช่องทางหารายได้เสริมที่มีความยืดหยุ่นและปลอดภัย เพราะอยากใช้เวลาว่างให้เกิดประโยชน์โดยไม่กระทบการเรียน'"
             ],
             "PROBLEM STATEMENTS", ACCENT_AMBER)

    # -------------------------------------------------------------
    # SLIDE 5: STEP 3 - IDEATE
    # -------------------------------------------------------------
    s5 = prs.slides.add_slide(blank_layout)
    add_header(s5, "ขั้นตอนที่ 3 · IDEATE (นาทีที่ 2:00 - 3:00)", "ระดมไอเดียแบบ Crazy 8 & คัดเลือกฟีเจอร์", "Divergent Thinking สู่ Convergent Thinking ด้วย Feature Prioritization Matrix")

    add_card(s5, Inches(0.8), Inches(1.9), Inches(4.5), Inches(4.9), "เทคนิค Crazy 8 (ระดม 8 ไอเดียใน 8 นาที)", 
             [
                 "ช่องที่ 1: ปุ่มกดเรียกรถแบบ One-touch ด่วน",
                 "ช่องที่ 2: แผนที่แสดงตำแหน่งรถแบบเรียลไทม์",
                 "ช่องที่ 3: ระบบแชร์รถทางเดียวกันหารค่าบริการ",
                 "ช่องที่ 4: ระบบจองคิวรถล่วงหน้าก่อนหมดคาบ",
                 "ช่องที่ 5: สแกนจ่ายด้วย PromptPay QR Code",
                 "ช่องที่ 6: ระบบให้คะแนนและรีวิวคนขับ",
                 "ช่องที่ 7: โหมดฉุกเฉินตอนฝนตกหนัก",
                 "ช่องที่ 8: กระดานจับคู่เพื่อนคณะเดียวกัน",
                 "กฎเหล็ก: เน้นปริมาณก่อนคุณภาพ ไม่ตัดสินไอเดีย"
             ],
             "DIVERGENT THINKING", ACCENT_PURPLE)

    add_card(s5, Inches(5.6), Inches(1.9), Inches(6.933), Inches(4.9), "Feature Prioritization Matrix (จัดลำดับฟีเจอร์สร้าง MVP)", 
             [
                 "MUST-HAVE (แกนหลักที่ต้องมีใน Prototype สองด้าน):",
                 "  1. [Consumer] ปุ่มเรียกรถและเลือกรอบบริการ — เพื่อตอบโจทย์ความเร่งรีบในการเดินทาง",
                 "  2. [Consumer] ตารางรอบรถว่างและราคาชัดเจน — ตัดสินใจได้ทันทีโดยไม่ต้องรอ",
                 "  3. [Provider] ฟอร์มสร้างรอบให้บริการ (เส้นทาง/ราคา) — ให้คนขับประกาศรอบว่างตามสะดวก",
                 "  4. [Both] ยืนยันตัวตนด้วยอีเมลมหาวิทยาลัย (.ac.th) — สร้าง Trust & Safety ระหว่างสองฝั่ง",
                 "----------------------------------------------------------------",
                 "NICE-TO-HAVE (เก็บไว้พัฒนาในเฟสถัดไปเพื่อไม่ให้ Scope บานปลาย):",
                 "  • [Consumer] ระบบแชร์รถหารค่าโดยสาร (GrabShare Style) — ดีไซน์ซับซ้อนเกินไปในรอบแรก",
                 "  • [Both] ระบบคะแนนรีวิวคนขับและผู้โดยสาร — สร้างความน่าเชื่อถือในระยะยาว"
             ],
             "FEATURE MATRIX & MVP SCOPING", ACCENT_BLUE)

    # -------------------------------------------------------------
    # SLIDE 6: STEP 4 - PROTOTYPE (PAPER & ARCHITECTURE)
    # -------------------------------------------------------------
    s6 = prs.slides.add_slide(blank_layout)
    add_header(s6, "ขั้นตอนที่ 4 · PROTOTYPE (รอบที่ 1: Paper สู่ สถาปัตยกรรม)", "จากภาพร่างกระดาษ สู่ระบบจริงด้วยพลัง AI", "ทดสอบ Low-fidelity เพื่อ Fail Cheap ก่อนเชื่อมโยง Google Sheet + Apps Script")

    add_card(s6, Inches(0.8), Inches(1.9), Inches(5.66), Inches(4.9), "รอบที่ 1: Paper Prototype ทดสอบ 2 หน้าจอ", 
             [
                 "การทดสอบความละเอียดต่ำ (Low-Fidelity):",
                 "  • ร่างหน้าจอ Provider: ฟอร์มระบุเส้นทาง, กำหนดราคา, และปุ่ม 'เปิดรับผู้โดยสาร'",
                 "  • ร่างหน้าจอ Consumer: รายการรอบรถที่ว่าง และปุ่ม 'จองเลย'",
                 "กระบวนการทดสอบด้วยกระดาษ:",
                 "  • ให้เพื่อน 5 คนใช้นิ้วจิ้มแทนเมาส์ (Paper Prototyping Test)",
                 "  • ข้อค้นพบ: ผู้ใช้สับสนว่าต้องกดยืนยันอีกครั้งหรือไม่",
                 "ข้อดีตามกฎเหล็ก 'Fail Fast, Fail Cheap':",
                 "  • แก้ไขตำแหน่งปุ่มและคำอธิบายบนกระดาษได้ใน 10 วินาที",
                 "  • ป้องกันการสั่ง AI เขียนโค้ดในสิ่งที่ไม่ตอบโจทย์ตั้งแต่ต้น"
             ],
             "LOW-FIDELITY PROTOTYPING", ACCENT_AMBER)

    add_card(s6, Inches(6.86), Inches(1.9), Inches(5.66), Inches(4.9), "สถาปัตยกรรมระบบ 3 ส่วน (Cloud Architecture)", 
             [
                 "1. Google Sheets (Database หลังบ้าน):",
                 "   - ชีต 'CampusRideData' จัดเก็บข้อมูลคอลัมน์:",
                 "     driverName | route | price | booked",
                 "2. Google Apps Script (Backend & Serverless API):",
                 "   - ไฟล์ Code.gs ทำหน้าที่อ่าน/เขียนข้อมูลใน Sheet",
                 "   - เสิร์ฟไฟล์หน้าเว็บ HTML สองฝั่งผ่านฟังก์ชัน doGet()",
                 "3. Gemini & Google Antigravity (AI Co-developer):",
                 "   - Agentic AI ช่วยวางแผนโครงสร้าง เขียนโค้ด Apps Script",
                 "   - ออกแบบ UI รองรับ Mobile-first ทั้งสองมุมมอง",
                 "ผลลัพธ์: Deploy เป็น Web App เข้าถึงได้ผ่าน URL บนมือถือทันที"
             ],
             "TECH STACK & ARCHITECTURE", ACCENT_GREEN)

    # -------------------------------------------------------------
    # SLIDE 7: STEP 4 - AI PROMPTS & CODE IMPLEMENTATION
    # -------------------------------------------------------------
    s7 = prs.slides.add_slide(blank_layout)
    add_header(s7, "ขั้นตอนที่ 4 · PROTOTYPE (การใช้ AI และโครงสร้างโค้ด)", "ชุดคำสั่ง Prompt สั่งการ AI สร้างแอปสองด้าน", "สร้างไฟล์ Code.gs, Provider.html และ Consumer.html โดยไม่ต้องเขียนโค้ดเอง")

    add_card(s7, Inches(0.8), Inches(1.9), Inches(5.66), Inches(4.9), "Prompt ฝั่งคนขับ (Provider.html)", 
             [
                 "พรอมต์สั่งงาน Gemini / Google Antigravity:",
                 "'ช่วยเขียนโค้ด Google Apps Script (Code.gs) และไฟล์ HTML ชื่อ Provider.html สำหรับเว็บแอปพลิเคชันที่เชื่อมกับ Google Sheet ชื่อ CampusRideData (คอลัมน์ driverName, route, price, booked)'",
                 "ข้อกำหนดฟังก์ชัน:",
                 "1. ฟอร์มกรอก ชื่อคนขับ, เส้นทาง, ราคา พร้อมปุ่ม 'เพิ่มรอบบริการ'",
                 "2. บันทึกข้อมูลขึ้น Google Sheet แบบ Real-time",
                 "3. แสดงตารางรอบบริการทั้งหมด พร้อมสถานะ 'ว่าง' หรือ 'มีผู้จองแล้ว'",
                 "4. ออกแบบ UI เรียบง่าย โทนสีฟ้า รองรับหน้าจอมือถือ (Responsive)",
                 "5. อธิบายวิธีการ Deploy เป็น Web App แบบ Anyone with link"
             ],
             "PROVIDER PROMPT & SPECS", ACCENT_BLUE)

    add_card(s7, Inches(6.86), Inches(1.9), Inches(5.66), Inches(4.9), "Prompt ฝั่งผู้โดยสาร (Consumer.html)", 
             [
                 "พรอมต์สั่งงาน Gemini / Google Antigravity:",
                 "'ช่วยเขียนไฟล์ HTML ชื่อ Consumer.html เพิ่มเติมจากโปรเจกต์เดิม โดยเชื่อมกับ Google Sheet เดียวกัน (CampusRideData)'",
                 "ข้อกำหนดฟังก์ชัน:",
                 "1. ดึงและแสดงเฉพาะรอบบริการที่มีสถานะ 'ว่าง' เท่านั้น",
                 "2. แสดงชื่อคนขับ เส้นทาง และราคาอย่างชัดเจน",
                 "3. มีปุ่ม 'จองเลย' ต่อท้ายแต่ละรายการ เมื่อกดแล้วให้อัปเดตสถานะในชีตเป็น 'มีผู้จองแล้ว' ทันที",
                 "4. แสดงข้อความแจ้งเตือนสีเขียวยืนยันการจองสำเร็จ",
                 "5. โทนสีฟ้าเข้ม สไตล์แอปพลิเคชันเรียกรถระดับสากล"
             ],
             "CONSUMER PROMPT & SPECS", ACCENT_CYAN)

    # -------------------------------------------------------------
    # SLIDE 8: LIVE DEMO SCREENSHOTS & WORKFLOW
    # -------------------------------------------------------------
    s8 = prs.slides.add_slide(blank_layout)
    add_header(s8, "การสาธิตผลงานจริง · LIVE DEMO (นาทีที่ 3:00 - 6:00)", "ทดสอบระบบจริง: การเชื่อมต่อของสองฝั่งผ่าน Cloud", "ข้อมูลไหลจากฝั่ง Provider สู่ Consumer และบันทึกลง Google Sheet แบบ Real-time")

    col4_w = Inches(2.75)
    add_card(s8, Inches(0.8), Inches(1.9), col4_w, Inches(4.9), "Step 1: คนขับเพิ่มรอบ", 
             [
                 "หน้า: Provider.html",
                 "การทำงาน:",
                 "• พี่บอลกรอกเส้นทาง:",
                 "  'ตึก 1 ไป โรงอาหาร 2'",
                 "• กำหนดราคา: 15 บาท",
                 "• กดปุ่ม 'เพิ่มรอบบริการ'",
                 "• ระบบส่งข้อมูลขึ้น Sheet ทันที"
             ],
             "PROVIDER ACTION", ACCENT_BLUE)

    add_card(s8, Inches(3.78), Inches(1.9), col4_w, Inches(4.9), "Step 2: ซิงค์ Cloud Database", 
             [
                 "ฐานข้อมูล: Google Sheet",
                 "การทำงาน:",
                 "• ตารางเพิ่มแถวใหม่อัตโนมัติ",
                 "• driverName: บอล",
                 "• route: ตึก 1 -> โรงอาหาร 2",
                 "• price: 15",
                 "• booked: 'ว่าง' (Available)",
                 "• ปลอดภัย ไม่ต้องเช่า Server"
             ],
             "CLOUD SYNC", ACCENT_GREEN)

    add_card(s8, Inches(6.76), Inches(1.9), col4_w, Inches(4.9), "Step 3: ผู้โดยสารกดจอง", 
             [
                 "หน้า: Consumer.html",
                 "การทำงาน:",
                 "• น้องปาล์มเปิดหน้ารายการ",
                 "• เลือกรอบของพี่บอล",
                 "• ตรวจสอบราคา 15 บาท",
                 "• กดปุ่ม 'จองเลย'",
                 "• ขึ้นแบนเนอร์เขียวยืนยันสำเร็จ"
             ],
             "CONSUMER ACTION", ACCENT_CYAN)

    add_card(s8, Inches(9.74), Inches(1.9), col4_w, Inches(4.9), "Step 4: อัปเดตสถานะ Realtime", 
             [
                 "ระบบ: Apps Script Trigger",
                 "การทำงาน:",
                 "• สถานะในชีตเปลี่ยนเป็น 'มีผู้จองแล้ว'",
                 "• รอบนี้หายจากหน้าค้นหา",
                 "• ป้องกันการกดจองซ้ำ",
                 "• ตอบโจทย์ Live Demo ครบทั้งสองฝั่ง (5 คะแนนเต็ม)"
             ],
             "SYSTEM UPDATE", ACCENT_AMBER)

    # -------------------------------------------------------------
    # SLIDE 9: STEP 5 - USABILITY TESTING & THINK ALOUD
    # -------------------------------------------------------------
    s9 = prs.slides.add_slide(blank_layout)
    add_header(s9, "ขั้นตอนที่ 5 · USABILITY TESTING (นาทีที่ 6:00 - 7:00)", "ผลการทดสอบกับผู้ใช้จริง 5 คน ด้วย Think Aloud", "แบ่งบทบาท Moderator, Note-taker, Observer พบปัญหาจริงเพื่อนำไปปรับปรุง")

    add_card(s9, Inches(0.8), Inches(1.9), Inches(5.66), Inches(4.9), "บันทึกผลการทดสอบผู้ใช้ (Usability Record)", 
             [
                 "ภารกิจ 1 (Provider): เพิ่มรอบให้บริการผ่าน provider.html",
                 "• ผล: สำเร็จ (ผู้ใช้พูด: 'อ๋อ กรอกเสร็จแล้วกดปุ่มนี้เลยใช่ไหม')",
                 "• ปัญหาที่พบ: ผู้ใช้ลังเลเล็กน้อยตรงตำแหน่งปุ่มส่งข้อมูล",
                 "----------------------------------------------------------------",
                 "ภารกิจ 2 (Consumer): ค้นหาและจองรถผ่าน consumer.html",
                 "• ผล: สำเร็จ (ผู้ใช้พูด: 'ดีนะ เห็นราคาก่อนกดจองเลย')",
                 "• ปัญหาที่พบ: อยากให้มีข้อความแจ้งเตือนสีเขียวชัดเจนขึ้นหลังจอง",
                 "----------------------------------------------------------------",
                 "ภารกิจ 3 (Provider): ตรวจสอบสถานะหลังมีคนจอง",
                 "• ผล: ไม่สำเร็จ (ผู้ใช้พูด: 'อ้าว ทำไมไม่เห็นอัปเดตว่าใครจอง')",
                 "• ปัญหาที่พบ: หน้าจอฝั่งคนขับไม่รีเฟรชสถานะแบบอัตโนมัติ"
             ],
             "FIELD USABILITY TESTING", ACCENT_BLUE)

    add_card(s9, Inches(6.86), Inches(1.9), Inches(5.66), Inches(4.9), "บทเรียนและการปรับปรุงแบบ Iterative Loop", 
             [
                 "ข้อค้นพบสำคัญ (Key Insight):",
                 "• หน้าจอฝั่ง Provider ขาดระบบแจ้งเตือนแบบทันที (Push Notification) ทำให้คนขับไม่ทราบว่ามีผู้โดยสารกดจองแล้ว",
                 "การดำเนินการแก้ไขอย่างรวดเร็ว (Pivot / Tweak):",
                 "• กลับไปสั่ง AI เพิ่มเติม: 'ช่วยเพิ่มการแจ้งเตือนแถบสีแดงและเสียงเตือนบนหน้า Provider เมื่อมีคนจองรอบแล้ว'",
                 "• AI แก้ไขโค้ดและส่งผลลัพธ์ใหม่ในเวลาไม่กี่นาที",
                 "คุณค่าของ Iteration:",
                 "• พิสูจน์หลักการ 'การย้อนกลับไม่ใช่ความล้มเหลว แต่คือ Learning Loop'",
                 "• นวัตกรรมสำเร็จด้วยการปรับปรุงจุดเล็ก ๆ ซ้ำแล้วซ้ำเล่า"
             ],
             "ITERATION & REFINEMENT", ACCENT_PURPLE)

    # -------------------------------------------------------------
    # SLIDE 10: STEP 5 - VIABILITY EXPERIMENTATION
    # -------------------------------------------------------------
    s10 = prs.slides.add_slide(blank_layout)
    add_header(s10, "ขั้นตอนที่ 5 · VIABILITY EXPERIMENT (ความยั่งยืนทางธุรกิจ)", "ทดสอบความต้องการจ่ายเงินจริงด้วย Fake Door Test", "หลักฐานเชิงประจักษ์ว่า 'มีคนยอมจ่ายเงินจริง' ก่อนลงทุนสร้างระบบเต็มรูปแบบ")

    add_card(s10, Inches(0.8), Inches(1.9), Inches(5.66), Inches(4.9), "การออกแบบการทดลอง Fake Door Test", 
             [
                 "สมมติฐานทางธุรกิจ (Business Hypothesis):",
                 "• นักศึกษายินดีจ่ายค่าบริการเสริมล่วงหน้าเพื่อการันตีว่าจะได้รถแน่นอน ไม่ต้องรอคิวในชั่วโมงเร่งด่วน",
                 "การสร้างประตูหลอก (Fake Door Setup):",
                 "• สร้างแบนเนอร์บนหน้าแรกและโพสต์ในกลุ่มมหาวิทยาลัย:",
                 "  'จองคิวรถรอบเช้าล่วงหน้า 5 บาท รับสิทธิ์ไม่ต้องรอคิวหน้างาน'",
                 "• ติดตั้งปุ่มวัดอัตราการคลิก (Call-to-Action Click)",
                 "เปรียบเทียบกับเทคนิคอื่น:",
                 "• ดีกว่าแบบสอบถาม เพราะแบบสอบถามวัดแค่ 'ความเห็น' แต่ Fake Door วัด 'พฤติกรรมจริง'"
             ],
             "EXPERIMENT SETUP", ACCENT_AMBER)

    add_card(s10, Inches(6.86), Inches(1.9), Inches(5.66), Inches(4.9), "ผลลัพธ์และตัวเลขชี้วัดความสำเร็จ", 
             [
                 "ตัวเลขผลลัพธ์จากการทดสอบ (Empirical Data):",
                 "• จำนวนผู้เห็นข้อความทดสอบทั้งหมด: 120 คน",
                 "• จำนวนผู้กดปุ่มแสดงความต้องการจองจริง: 43 คน",
                 "• Conversion Rate สูงถึง: 35.8% (~36%)",
                 "การวิเคราะห์ความคุ้มค่าและทิศทางธุรกิจ:",
                 "• ได้รับสัญญาณเชิงบวกชัดเจน (Strong Buying Signal)",
                 "• ประมาณการรายได้: หากมีผู้ใช้งาน 300 เที่ยว/วัน หักค่า GP 10% แพลตฟอร์มจะมีรายได้ 450 บาท/วัน (คุ้มค่าดูแลระบบ)",
                 "สรุป: ไอเดียผ่านการตรวจสอบครบทั้งสามเสาหลัก",
                 "  [✓] Desirability   [✓] Feasibility   [✓] Viability"
             ],
             "METRICS & SIGNAL VALIDATION", ACCENT_GREEN)

    # -------------------------------------------------------------
    # SLIDE 11: DECISION LOGIC & LESSONS LEARNED
    # -------------------------------------------------------------
    s11 = prs.slides.add_slide(blank_layout)
    add_header(s11, "บทสรุป · DECISION LOGIC & NEXT STEPS", "แผนผังการตัดสินใจเพื่อก้าวต่อไปของนวัตกรรม", "ผลรวมของการปรับปรุงเล็ก ๆ น้อย ๆ และก้าวต่อไปสู่การ Scale Up")

    col3_w = Inches(3.64)
    add_card(s11, Inches(0.8), Inches(1.9), col3_w, Inches(4.9), "ระดับ 1: Desirability?", 
             [
                 "คำถามตรวจสอบ:",
                 "แก้ปัญหาได้จริงไหม?",
                 "-----------------------",
                 "• ผล: [ผ่านเกณฑ์]",
                 "• นักศึกษา 100% ตอบรับว่าลดเวลาเดินข้ามตึกจาก 25 นาทีเหลือ 5 นาที",
                 "• ถ้าไม่ผ่าน: ต้องหยุดและย้อนกลับไป Define ปัญหาใหม่"
             ],
             "STAGE 1: PROBLEM-SOLUTION FIT", ACCENT_BLUE)

    add_card(s11, Inches(4.84), Inches(1.9), col3_w, Inches(4.9), "ระดับ 2: Feasibility?", 
             [
                 "คำถามตรวจสอบ:",
                 "สร้างได้และใช้งานง่ายไหม?",
                 "-----------------------",
                 "• ผล: [ผ่านเกณฑ์หลัง Tweak]",
                 "• ใช้งาน Google Apps Script + Sheets สำเร็จ 100%",
                 "• ปรับแก้ตาม Usability Test: เพิ่มแจ้งเตือนคนขับและปุ่มที่ชัดเจนขึ้น"
             ],
             "STAGE 2: TECHNICAL VIABILITY", ACCENT_GREEN)

    add_card(s11, Inches(8.88), Inches(1.9), col3_w, Inches(4.9), "ระดับ 3: Viability?", 
             [
                 "คำถามตรวจสอบ:",
                 "มีสัญญาณคนยอมจ่ายเงินไหม?",
                 "-----------------------",
                 "• ผล: [ผ่านเกณฑ์]",
                 "• Fake Door Test ได้ผลตอบรับ 36% แสดงถึงความเต็มใจจ่าย",
                 "• ก้าวต่อไป (Scale Up): ขยายจุดจอดรับส่ง และเชื่อมต่อระบบ PromptPay จริง"
             ],
             "STAGE 3: COMMERCIAL VIABILITY", ACCENT_AMBER)

    # -------------------------------------------------------------
    # SLIDE 12: CONCLUSION & Q&A
    # -------------------------------------------------------------
    s12 = prs.slides.add_slide(blank_layout)
    add_header(s12, "บทสรุป · Q&A SESSION", "สรุปบทเรียนและเปิดรับคำถาม (Ready for Q&A)", "การประเมินผลตามเกณฑ์ Rubric Scoring เต็ม 100 คะแนน และ Pitching 20 คะแนน")

    add_card(s12, Inches(0.8), Inches(1.9), Inches(5.66), Inches(4.9), "สรุปบทเรียนสำคัญจาก Playbook", 
             [
                 "1. 'Empathy First': การรับฟังโดยไม่ตัดสินคือก้าวแรกของนวัตกรรมที่แท้จริง",
                 "2. 'AI as a Co-creator': ไม่ต้องเขียนโค้ดเป็น แต่ต้องเขียน Prompt สื่อสารกับ AI ให้ชัดเจน",
                 "3. 'Two-sided Alignment': แพลตฟอร์มสองด้านต้องสร้างคุณค่าให้ทั้งสองฝั่งเท่าเทียมกัน",
                 "4. 'Fail Fast, Fail Cheap': การทดสอบ Usability และ Fake Door ช่วยลดความเสี่ยงทางธุรกิจได้อย่างมหาศาล",
                 "5. นวัตกรรมไม่ได้จบที่การนำเสนอ แต่คือวงจรเรียนรู้ที่ไม่มีวันสิ้นสุด"
             ],
             "KEY TAKEAWAYS", ACCENT_PURPLE)

    add_card(s12, Inches(6.86), Inches(1.9), Inches(5.66), Inches(4.9), "ความพร้อมในการตอบข้อซักถาม (Q&A)", 
             [
                 "ประเด็นที่เตรียมพร้อมตอบข้อซักถาม:",
                 "• ความปลอดภัยและความน่าเชื่อถือของผู้ขับขี่ (Safety Verification)",
                 "• ความเสถียรของสถาปัตยกรรม Google Apps Script เมื่อผู้ใช้เพิ่มขึ้น",
                 "• โมเดลรายได้และโครงสร้างการแบ่งรายได้คนขับ (Revenue Sharing Model)",
                 "• ผลการทำ Usability Testing และ Iteration แผนงานในอนาคต",
                 "----------------------------------------------------------------",
                 "ขอขอบคุณคณาจารย์และเพื่อนร่วมชั้นทุกท่าน",
                 "คณะครุศาสตร์อุตสาหกรรมและเทคโนโลยี มทร.ศรีวิชัย"
             ],
             "OPEN FOR QUESTIONS", ACCENT_GREEN)

    out_path = r"c:\Users\pasit\OneDrive\Desktop\ENERGY SERVE\Design_Thinking_Playbook_Presentation.pptx"
    prs.save(out_path)
    print(f"Presentation saved successfully to: {out_path}")

if __name__ == "__main__":
    create_deck()
